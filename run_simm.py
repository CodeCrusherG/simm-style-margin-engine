"""Run: sensitivities -> SIMM-style IM -> margin back-test -> Excel + plot."""
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

import im_backtest as bt
from pricers import base_market, demo_portfolio, portfolio_pv
from sensitivities import portfolio_sensitivities
from simm import simm_im


def main():
    m, trades = base_market(), demo_portfolio()
    print(f"Portfolio PV: {portfolio_pv(trades, m):,.0f}")

    sens = portfolio_sensitivities(trades, m)
    total, breakdown = simm_im(sens)
    print(f"\nSIMM-style IM (illustrative parameters): {total:,.0f}")
    bd = pd.DataFrame(breakdown).T.round(0)
    print(bd, "\n")

    f = bt.simulate_factors()
    w = bt.window_sums(f)
    pnl = bt.window_pnl(trades, m, w)
    result, hvar_im = bt.backtest_im(pnl, total)
    print(result.round(3).to_string(index=False))

    out = "simm_report.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        sens.round(2).to_excel(xw, sheet_name="Sensitivities", index=False)
        bd.to_excel(xw, sheet_name="IM_breakdown")
        result.round(3).to_excel(xw, sheet_name="IM_backtest", index=False)
        pd.DataFrame({"window_pnl": pnl}).round(0).to_excel(
            xw, sheet_name="Window_PnL", index=False)
    wb = load_workbook(out)
    for ws in wb:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 24
    wb.save(out)

    k = int(len(pnl) * 0.6)
    plt.figure(figsize=(10, 5))
    plt.plot(pnl, lw=0.7, label="10-day P&L (full revaluation)")
    plt.axhline(-total, color="red", label="-SIMM-style IM")
    plt.axhline(-hvar_im, color="orange", ls="--", label="-Historical-VaR IM")
    plt.axvline(k, color="grey", ls=":", label="calibration | test")
    plt.legend(); plt.xlabel("10-day window"); plt.ylabel("P&L")
    plt.title("Initial margin vs realised 10-day losses (simulated factors)")
    plt.tight_layout(); plt.savefig("im_backtest_plot.png", dpi=130)


if __name__ == "__main__":
    main()
